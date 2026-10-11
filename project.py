"""Simple Excel chart viewer for the CS50P final project."""

import re
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook


CHART_TYPES = ("Scatter", "Line")
NUMBER_PATTERN = r"[-+]?(?:\d{1,3}(?:\.\d{3})+|\d+(?:,\d+)?)"
TIME_PATTERN = re.compile(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$")
TIME_GROUPS = {
    "No grouping": 0,
    "30 seconds": 30,
    "1 minute": 60,
    "5 minutes": 300,
    "15 minutes": 900,
    "1 hour": 3600,
}
TIME_PLACEHOLDER = "HH:MM:SS"


def parse_number(value):
    """Read the first number from the spreadsheet's fixed value format."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value

    match = re.search(NUMBER_PATTERN, str(value).strip())
    if not match:
        return None

    result_text = match.group()
    sign = ""
    if result_text.startswith(("+", "-")):
        sign, result_text = result_text[0], result_text[1:]
    if "." in result_text and "," not in result_text:
        number = result_text.replace(".", "")
    else:
        number = result_text.replace(",", ".")
    try:
        return float(sign + number)
    except ValueError:
        return None


def format_axis_number(value, _position):
    """Show large chart values with dots as thousands separators."""
    if abs(value) < 1000:
        text = f"{value:,.2f}".rstrip("0").rstrip(".")
    else:
        text = f"{value:,.0f}"
    return text.replace(",", ".")


def parse_time_value(value):
    """Convert HH:MM[:SS] text to seconds since midnight for charting."""
    match = TIME_PATTERN.fullmatch(str(value).strip())
    if not match:
        return None
    hour, minute, second = (int(part or 0) for part in match.groups())
    return hour * 3600 + minute * 60 + second


def is_time_column(rows, column_index):
    """Return whether the non-empty values in a column are times."""
    values = [row[column_index] for row in rows
              if column_index < len(row) and row[column_index] not in (None, "")]
    return bool(values) and all(parse_time_value(value) is not None for value in values)


def format_time_axis(value, _position):
    """Format chart time ticks as HH:MM:SS."""
    total_seconds = max(0, int(round(value)))
    hour, remainder = divmod(total_seconds, 3600)
    minute, second = divmod(remainder, 60)
    return f"{hour:02d}:{minute:02d}:{second:02d}"


def format_hover_value(value, is_time=False):
    """Format a chart value for the mouse hover tooltip."""
    if is_time:
        return format_time_axis(value, 0)
    if isinstance(value, float) and value.is_integer():
        return f"{value:.0f}"
    return format_axis_number(value, 0) if isinstance(value, (int, float)) else str(value)


def add_hover_tooltip(figure, axis, points, x_label, y_label,
                      x_is_time, y_is_time):
    """Show the closest data point while the mouse moves over the chart."""
    annotation = axis.annotate(
        "", xy=(0, 0), xytext=(12, 12), textcoords="offset points",
        bbox={"boxstyle": "round", "fc": "white", "ec": "#264653", "alpha": 0.9},
    )
    annotation.set_visible(False)

    def on_move(event):
        if event.inaxes != axis or event.x is None or event.y is None:
            annotation.set_visible(False)
            figure.canvas.draw_idle()
            return

        closest = None
        closest_distance = 10
        for x_value, y_value in points:
            if not isinstance(x_value, (int, float)):
                continue
            pixel_x, pixel_y = axis.transData.transform((x_value, y_value))
            distance = ((pixel_x - event.x) ** 2 + (pixel_y - event.y) ** 2) ** 0.5
            if distance < closest_distance:
                closest = (x_value, y_value)
                closest_distance = distance

        if closest is None:
            annotation.set_visible(False)
        else:
            x_value, y_value = closest
            annotation.xy = (x_value, y_value)
            annotation.set_text(
                f"{x_label}: {format_hover_value(x_value, x_is_time)}\n"
                f"{y_label}: {format_hover_value(y_value, y_is_time)}"
            )
            annotation.set_visible(True)
        figure.canvas.draw_idle()

    figure.canvas.mpl_connect("motion_notify_event", on_move)


def aggregate_by_time(x_values, y_values, interval_seconds):
    """Group time-series rows and keep the last Y value in each time bucket."""
    if interval_seconds == 0:
        return list(x_values), list(y_values)
    grouped = {}
    for x_value, y_value in zip(x_values, y_values):
        seconds = parse_time_value(x_value)
        if seconds is None:
            raise ValueError("Grouping requires a time column on the X axis")
        bucket = seconds // interval_seconds * interval_seconds
        grouped[bucket] = (format_time_axis(bucket, 0), y_value)
    ordered = [grouped[bucket] for bucket in sorted(grouped)]
    return ([value[0] for value in ordered],
            [value[1] for value in ordered])


def filter_by_time(x_values, y_values, start_text, end_text, time_on_x):
    """Keep rows inside an optional start/end time range."""
    start = parse_time_value(start_text) if start_text.strip() else None
    end = parse_time_value(end_text) if end_text.strip() else None
    if ((start_text.strip() and start is None)
            or (end_text.strip() and end is None)):
        raise ValueError("Time range must use HH:MM or HH:MM:SS")
    if start is not None and end is not None and start > end:
        raise ValueError("Start time must be before end time")

    time_values = x_values if time_on_x else y_values
    filtered_x = []
    filtered_y = []
    for x_value, y_value, time_value in zip(x_values, y_values, time_values):
        seconds = parse_time_value(time_value)
        if seconds is None:
            continue
        if start is not None and seconds < start:
            continue
        if end is not None and seconds > end:
            continue
        filtered_x.append(x_value)
        filtered_y.append(y_value)
    return filtered_x, filtered_y


def load_excel(path):
    """Read the first worksheet and return headers and non-empty data rows."""
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value) if value is not None else f"Column {index + 1}"
                   for index, value in enumerate(next(rows))]
        data = [list(row) for row in rows if any(value is not None for value in row)]
        return headers, data
    finally:
        workbook.close()


def create_chart(x_values, y_values, chart_type, x_label="X", y_label="Y"):
    """Build and return a matplotlib figure for the selected values."""
    if chart_type not in CHART_TYPES:
        raise ValueError("Unknown chart type")
    if not y_values:
        raise ValueError("There is no data to draw")

    time_values = [parse_time_value(value) for value in x_values]
    is_time_axis = bool(x_values) and all(value is not None for value in time_values)
    if is_time_axis:
        x_values = time_values
    else:
        x_values = [parse_number(value) for value in x_values]
    time_y_values = [parse_time_value(value) for value in y_values]
    is_time_y_axis = bool(y_values) and all(value is not None for value in time_y_values)
    if is_time_y_axis:
        y_values = time_y_values
    else:
        y_values = [parse_number(value) for value in y_values]

    figure, axis = plt.subplots(figsize=(8, 5))
    pairs = [(x_value, y_value) for x_value, y_value in zip(x_values, y_values)
             if isinstance(y_value, (int, float))]
    if not pairs:
        raise ValueError(f"{chart_type} chart needs a numeric Y column")

    if chart_type == "Line":
        axis.plot([x_value for x_value, _ in pairs],
                  [y_value for _, y_value in pairs], linewidth=0.8, color="#264653")
    else:
        numeric_x = [x_value if isinstance(x_value, (int, float)) else index
                     for index, (x_value, _) in enumerate(pairs)]
        axis.scatter(numeric_x, [y_value for _, y_value in pairs],
                     alpha=0.7, color="#e76f51")
        unique_x = sorted(set(value for value in numeric_x if isinstance(value, (int, float))))
        if len(unique_x) <= 30:
            axis.set_xticks(unique_x)
            axis.xaxis.set_major_formatter(FuncFormatter(
                lambda value, _position: f"{value:g}"))

    add_hover_tooltip(figure, axis, pairs, x_label, y_label,
                      is_time_axis, is_time_y_axis)

    axis.set_xlabel(x_label)
    axis.set_ylabel(y_label)
    if is_time_axis:
        axis.xaxis.set_major_formatter(FuncFormatter(format_time_axis))
    if is_time_y_axis:
        axis.yaxis.set_major_formatter(FuncFormatter(format_time_axis))
    else:
        axis.yaxis.set_major_formatter(FuncFormatter(format_axis_number))
    axis.set_title(f"{chart_type} chart")
    figure.tight_layout()
    return figure


class ExcelChartApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel Chart Viewer")
        self.root.geometry("1050x700")
        self.headers = []
        self.rows = []
        self.figure = None
        self.canvas = None
        self.toolbar = None

        controls = ttk.Frame(root, padding=10)
        controls.pack(fill="x")
        ttk.Button(controls, text="Open Excel file", command=self.open_file).pack(side="left")
        self.file_label = ttk.Label(controls, text="No file selected")
        self.file_label.pack(side="left", padx=10)
        ttk.Label(controls, text="Chart:").pack(side="left", padx=(20, 4))
        self.chart_choice = ttk.Combobox(controls, values=CHART_TYPES, state="readonly", width=12)
        self.chart_choice.set("Scatter")
        self.chart_choice.pack(side="left")
        ttk.Label(controls, text="X column:").pack(side="left", padx=(12, 4))
        self.x_choice = ttk.Combobox(controls, state="readonly", width=18)
        self.x_choice.pack(side="left")
        self.x_choice.bind("<<ComboboxSelected>>", self.update_group_control)
        ttk.Label(controls, text="Y column:").pack(side="left", padx=(8, 4))
        self.y_choice = ttk.Combobox(controls, state="readonly", width=18)
        self.y_choice.pack(side="left")
        self.y_choice.bind("<<ComboboxSelected>>", self.update_group_control)
        ttk.Label(controls, text="Group:").pack(side="left", padx=(8, 4))
        self.group_choice = ttk.Combobox(controls, values=tuple(TIME_GROUPS),
                         state="readonly", width=12)
        self.group_choice.set("No grouping")
        self.group_choice.configure(state="disabled")
        self.group_choice.pack(side="left")
        ttk.Label(controls, text="From:").pack(side="left", padx=(8, 4))
        self.start_time = ttk.Entry(controls, width=10)
        self.start_time.pack(side="left")
        ttk.Label(controls, text="To:").pack(side="left", padx=(4, 4))
        self.end_time = ttk.Entry(controls, width=10)
        self.end_time.pack(side="left")
        self._set_time_placeholder(self.start_time)
        self._set_time_placeholder(self.end_time)
        self.start_time.bind("<FocusIn>", self._clear_time_placeholder)
        self.start_time.bind("<FocusOut>", self._restore_time_placeholder)
        self.end_time.bind("<FocusIn>", self._clear_time_placeholder)
        self.end_time.bind("<FocusOut>", self._restore_time_placeholder)
        self.start_time.configure(state="disabled")
        self.end_time.configure(state="disabled")
        ttk.Button(controls, text="Draw chart", command=self.draw_chart).pack(side="left", padx=8)

        self.selection_hint = ttk.Label(root, text="Choose X and Y columns to draw a chart.")
        self.selection_hint.pack(anchor="w", padx=12)

        self.table = ttk.Treeview(root, show="headings", height=8)
        self.table.pack(fill="x", padx=12)
        self.chart_area = ttk.Frame(root)
        self.chart_area.pack(fill="both", expand=True, padx=12, pady=10)

    def open_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xlsm")])
        if not path:
            return
        try:
            self.headers, self.rows = load_excel(path)
        except Exception as error:
            messagebox.showerror("Cannot open file", str(error))
            return

        self.file_label.config(text=Path(path).name)
        self.x_choice["values"] = self.headers
        self.y_choice["values"] = self.headers
        if self.headers:
            self.x_choice.current(0)
            self.y_choice.current(1 if len(self.headers) > 1 else 0)
        self.show_preview()
        self.update_group_control()

    def _set_time_placeholder(self, entry):
        """Show an inactive example in an empty time entry."""
        entry.delete(0, tk.END)
        entry.insert(0, TIME_PLACEHOLDER)
        entry.configure(foreground="gray")

    def _clear_time_placeholder(self, event):
        if event.widget.get() == TIME_PLACEHOLDER:
            event.widget.delete(0, tk.END)
            event.widget.configure(foreground="black")

    def _restore_time_placeholder(self, event):
        if not event.widget.get().strip():
            self._set_time_placeholder(event.widget)

    def _time_entry_value(self, entry):
        value = entry.get().strip()
        return "" if value == TIME_PLACEHOLDER else value

    def update_group_control(self, _event=None):
        """Enable grouping only when X or Y contains time values."""
        has_time = False
        for choice in (self.x_choice, self.y_choice):
            if choice.get() and choice.get() in self.headers:
                has_time = has_time or is_time_column(
                    self.rows, self.headers.index(choice.get()))
        if has_time:
            self.group_choice.configure(state="readonly")
            self.start_time.configure(state="normal")
            self.end_time.configure(state="normal")
        else:
            self.group_choice.set("No grouping")
            self.group_choice.configure(state="disabled")
            self._set_time_placeholder(self.start_time)
            self._set_time_placeholder(self.end_time)
            self.start_time.configure(state="disabled")
            self.end_time.configure(state="disabled")

    def show_preview(self):
        self.table.delete(*self.table.get_children())
        self.table["columns"] = self.headers
        for header in self.headers:
            self.table.heading(header, text=header)
            self.table.column(header, width=150)
        for row in self.rows[:10]:
            self.table.insert("", tk.END, values=row)

    def draw_chart(self):
        chart_type = self.chart_choice.get()
        if not self.headers or not self.y_choice.get():
            messagebox.showwarning("Choose columns", "Open a file and choose a Y column.")
            return

        y_index = self.headers.index(self.y_choice.get())
        y_values = [row[y_index] for row in self.rows]
        try:
            if not self.x_choice.get():
                messagebox.showwarning("Choose columns", "Choose an X column.")
                return
            x_index = self.headers.index(self.x_choice.get())
            if x_index == y_index:
                messagebox.showwarning("Choose columns", "X and Y columns must be different.")
                return
            x_values = [row[x_index] for row in self.rows]
            time_on_x = is_time_column(self.rows, x_index)
            time_on_y = is_time_column(self.rows, y_index)
            if time_on_x or time_on_y:
                x_values, y_values = filter_by_time(
                    x_values, y_values,
                    self._time_entry_value(self.start_time),
                    self._time_entry_value(self.end_time), time_on_x)
            interval_seconds = TIME_GROUPS[self.group_choice.get()]
            if interval_seconds:
                if time_on_x:
                    x_values, y_values = aggregate_by_time(x_values, y_values, interval_seconds)
                elif time_on_y:
                    y_values, x_values = aggregate_by_time(y_values, x_values, interval_seconds)
            figure = create_chart(x_values, y_values, chart_type,
                                  self.x_choice.get(), self.y_choice.get())
        except (TypeError, ValueError) as error:
            messagebox.showerror("Cannot draw chart", str(error))
            return

        if self.canvas is not None:
            self.canvas.get_tk_widget().destroy()
        if self.toolbar is not None:
            self.toolbar.destroy()
        if self.figure is not None:
            plt.close(self.figure)
        self.figure = figure
        self.canvas = FigureCanvasTkAgg(figure, master=self.chart_area)
        self.canvas.draw()
        self.canvas.get_tk_widget().pack(fill="both", expand=True)
        self.toolbar = NavigationToolbar2Tk(self.canvas, self.chart_area, pack_toolbar=False)
        self.toolbar.update()
        self.toolbar.pack(fill="x")
        self.toolbar.pan()

def main():
    root = tk.Tk()
    ExcelChartApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()